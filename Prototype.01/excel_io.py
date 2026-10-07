# excel_io.py
"""Экспорт и импорт данных в/из Excel (.xlsx) через openpyxl.

Формат файла — один лист на сущность:
  Группы, Аудитории, Преподаватели, Предметы,
  Преподаватель-Предмет, Преподаватель-Аудитория, Курс-Предмет,
  Расписание.

Импорт идемпотентен: записи сопоставляются по названиям,
существующие обновляются, отсутствующие создаются.
"""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from db import KIND_LECTURE, KIND_PRACTICE

HEADER_FILL = PatternFill(start_color="3D3D3D", end_color="3D3D3D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="E6E6E6")

KIND_TO_RU = {KIND_LECTURE: "Лекция", KIND_PRACTICE: "Практика"}
RU_TO_KIND = {
    "лекция": KIND_LECTURE, "lecture": KIND_LECTURE, "л": KIND_LECTURE,
    "практика": KIND_PRACTICE, "practice": KIND_PRACTICE, "п": KIND_PRACTICE,
}
WEEKDAYS_RU = ["Пн", "Вт", "Ср", "Чт", "Пт"]

SHEETS = {
    "groups": ("Группы", ["Название", "Студентов", "Курс"]),
    "auditoriums": ("Аудитории", ["Название", "Вместимость", "Тип"]),
    "teachers": ("Преподаватели", ["Имя"]),
    "subjects": ("Предметы", ["Название"]),
    "teacher_subjects": ("Преподаватель-Предмет", ["Преподаватель", "Предмет"]),
    "teacher_auditoriums": ("Преподаватель-Аудитория", ["Преподаватель", "Аудитория"]),
    "course_subjects": ("Курс-Предмет", ["Курс", "Предмет"]),
    "schedule": ("Расписание", ["Преподаватель", "Предмет", "Группа",
                                "Аудитория", "Тип", "Пара", "День"]),
}


def _write_sheet(wb, key, rows):
    title, headers = SHEETS[key]
    ws = wb.create_sheet(title)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for r, row in enumerate(rows, 2):
        for col, value in enumerate(row, 1):
            ws.cell(row=r, column=col, value=value)
    for col in range(1, len(headers) + 1):
        widths = [len(str(ws.cell(row=r, column=col).value or ""))
                  for r in range(1, ws.max_row + 1)]
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = (
            min(max(max(widths, default=10) + 2, 12), 50)
        )


def export_to_excel(db, path: str) -> dict:
    """Выгружает все данные БД в .xlsx. Возвращает {лист: кол-во строк}."""
    wb = Workbook()
    wb.remove(wb.active)
    stats = {}

    groups = db.get_groups()
    _write_sheet(wb, "groups",
                 [(g["name"], g["students"], g["course"]) for g in groups])
    stats["Группы"] = len(groups)

    auds = db.get_auditoriums()
    _write_sheet(wb, "auditoriums",
                 [(a.name, a.capacity, KIND_TO_RU.get(a.kind, a.kind)) for a in auds])
    stats["Аудитории"] = len(auds)

    teachers = db.get_teachers()
    _write_sheet(wb, "teachers", [(t["name"],) for t in teachers])
    stats["Преподаватели"] = len(teachers)

    subjects = db.get_subjects()
    _write_sheet(wb, "subjects", [(s["name"],) for s in subjects])
    stats["Предметы"] = len(subjects)

    ts_rows = []
    for t in teachers:
        for s in db.get_subjects_of_teacher(int(t["id"])):
            ts_rows.append((t["name"], s["name"]))
    _write_sheet(wb, "teacher_subjects", ts_rows)

    ta_rows = []
    for t in teachers:
        for a in db.get_auditoriums_of_teacher(int(t["id"])):
            ta_rows.append((t["name"], a.name))
    _write_sheet(wb, "teacher_auditoriums", ta_rows)

    cs_rows = []
    for s in subjects:
        for course in db.get_courses_of_subject(int(s["id"])):
            cs_rows.append((course, s["name"]))
    _write_sheet(wb, "course_subjects", cs_rows)

    sched_rows = []
    for t in teachers:
        for r in db.get_teacher_schedule(int(t["id"])):
            sched_rows.append((
                t["name"],
                r.get("subject_name") or "",
                r.get("group_name", ""),
                r.get("aud_name", ""),
                KIND_TO_RU.get(r.get("kind"), r.get("kind")),
                r.get("pair_number", 1),
                WEEKDAYS_RU[int(r.get("weekday", 0) or 0)]
                if 0 <= int(r.get("weekday", 0) or 0) <= 4 else 0,
            ))
    _write_sheet(wb, "schedule", sched_rows)
    stats["Расписание"] = len(sched_rows)

    wb.save(path)
    return stats


def _read_sheet(wb, key):
    """Возвращает список словарей {заголовок: значение} (пустые строки пропуск)."""
    title, _headers = SHEETS[key]
    if title not in wb.sheetnames:
        return []
    ws = wb[title]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [(str(h).strip() if h is not None else "") for h in rows[0]]
    out = []
    for values in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        out.append({h: (values[i] if i < len(values) else None)
                    for i, h in enumerate(headers)})
    return out


def _parse_kind(value) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    kind = RU_TO_KIND.get(str(value).strip().lower())
    if kind is None:
        raise ValueError(f"неизвестный тип: {value!r} (нужно Лекция/Практика)")
    return kind


def _parse_weekday(value) -> int:
    if value is None or str(value).strip() == "":
        return 0
    s = str(value).strip()
    if s.isdigit() and 0 <= int(s) <= 4:
        return int(s)
    low = s.lower()
    for i, name in enumerate(["пн", "вт", "ср", "чт", "пт"]):
        if low.startswith(name):
            return i
    raise ValueError(f"неизвестный день: {value!r} (нужно Пн..Пт или 0..4)")


def _find_id(db, table: str, name: str):
    """Ищет id записи по точному названию. None если нет."""
    row = db.c.execute(f"SELECT id FROM {table} WHERE name = ?", (name,)).fetchone()
    return int(row["id"]) if row else None


def import_from_excel(db, path: str) -> dict:
    """Загружает данные из .xlsx в БД.

    Возвращает {"added": {...}, "updated": {...}, "errors": [...]}.
    Ошибки в отдельных строках не прерывают импорт остальных.
    """
    wb = load_workbook(path, data_only=True)
    added: dict[str, int] = {}
    updated: dict[str, int] = {}
    errors: list[str] = []

    def bump(d, k):
        d[k] = d.get(k, 0) + 1

    # --- Группы ---
    for i, row in enumerate(_read_sheet(wb, "groups"), 2):
        try:
            name = str(row.get("Название", "") or "").strip()
            if not name:
                raise ValueError("пустое название")
            students = int(row.get("Студентов", 0))
            course = int(row.get("Курс", 0))
            if students <= 0 or not 1 <= course <= 6:
                raise ValueError(f"студентов={row.get('Студентов')}, курс={row.get('Курс')}")
            gid = _find_id(db, "groups", name)
            if gid is None:
                db.add_group(name, students, course)
                bump(added, "Группы")
            else:
                db.update_group(gid, name, students, course)
                bump(updated, "Группы")
        except Exception as e:
            errors.append(f"Группы, строка {i}: {e}")

    # --- Аудитории ---
    for i, row in enumerate(_read_sheet(wb, "auditoriums"), 2):
        try:
            name = str(row.get("Название", "") or "").strip()
            if not name:
                raise ValueError("пустое название")
            capacity = int(row.get("Вместимость", 0))
            kind = _parse_kind(row.get("Тип")) or KIND_LECTURE
            if capacity <= 0:
                raise ValueError(f"вместимость={row.get('Вместимость')}")
            aid = _find_id(db, "auditoriums", name)
            if aid is None:
                db.add_auditorium(name, capacity, kind)
                bump(added, "Аудитории")
            else:
                db.update_auditorium(aid, name, capacity)
                db.update_auditorium_kind(aid, kind)
                bump(updated, "Аудитории")
        except Exception as e:
            errors.append(f"Аудитории, строка {i}: {e}")

    # --- Преподаватели ---
    for i, row in enumerate(_read_sheet(wb, "teachers"), 2):
        try:
            name = str(row.get("Имя", "") or "").strip()
            if not name:
                raise ValueError("пустое имя")
            if _find_id(db, "teachers", name) is None:
                db.add_teacher(name)
                bump(added, "Преподаватели")
        except Exception as e:
            errors.append(f"Преподаватели, строка {i}: {e}")

    # --- Предметы ---
    for i, row in enumerate(_read_sheet(wb, "subjects"), 2):
        try:
            name = str(row.get("Название", "") or "").strip()
            if not name:
                raise ValueError("пустое название")
            if _find_id(db, "subjects", name) is None:
                db.add_subject(name)
                bump(added, "Предметы")
        except Exception as e:
            errors.append(f"Предметы, строка {i}: {e}")

    # --- Связка преподаватель-предмет ---
    for i, row in enumerate(_read_sheet(wb, "teacher_subjects"), 2):
        try:
            t = str(row.get("Преподаватель", "") or "").strip()
            s = str(row.get("Предмет", "") or "").strip()
            tid, sid = _find_id(db, "teachers", t), _find_id(db, "subjects", s)
            if tid is None or sid is None:
                raise ValueError(f"не найден: преподаватель={t!r} / предмет={s!r}")
            db.assign_teacher_subject(tid, sid)
            bump(added, "Связки препод-предмет")
        except Exception as e:
            errors.append(f"Преподаватель-Предмет, строка {i}: {e}")

    # --- Связка преподаватель-аудитория ---
    for i, row in enumerate(_read_sheet(wb, "teacher_auditoriums"), 2):
        try:
            t = str(row.get("Преподаватель", "") or "").strip()
            a = str(row.get("Аудитория", "") or "").strip()
            tid, aid = _find_id(db, "teachers", t), _find_id(db, "auditoriums", a)
            if tid is None or aid is None:
                raise ValueError(f"не найден: преподаватель={t!r} / аудитория={a!r}")
            db.assign_teacher_auditorium(tid, aid)
            bump(added, "Связки препод-аудитория")
        except Exception as e:
            errors.append(f"Преподаватель-Аудитория, строка {i}: {e}")

    # --- Связка курс-предмет ---
    for i, row in enumerate(_read_sheet(wb, "course_subjects"), 2):
        try:
            course = int(row.get("Курс", 0))
            s = str(row.get("Предмет", "") or "").strip()
            sid = _find_id(db, "subjects", s)
            if not 1 <= course <= 6 or sid is None:
                raise ValueError(f"курс={row.get('Курс')}, предмет={s!r}")
            db.add_course_subject(course, sid)
            bump(added, "Связки курс-предмет")
        except Exception as e:
            errors.append(f"Курс-Предмет, строка {i}: {e}")

    # --- Расписание ---
    for i, row in enumerate(_read_sheet(wb, "schedule"), 2):
        try:
            t = str(row.get("Преподаватель", "") or "").strip()
            s = str(row.get("Предмет", "") or "").strip()
            g = str(row.get("Группа", "") or "").strip()
            a = str(row.get("Аудитория", "") or "").strip()
            tid = _find_id(db, "teachers", t)
            gid = _find_id(db, "groups", g)
            aid = _find_id(db, "auditoriums", a)
            sid = _find_id(db, "subjects", s) if s else None
            kind = _parse_kind(row.get("Тип")) or KIND_LECTURE
            pair = int(row.get("Пара", 1))
            weekday = _parse_weekday(row.get("День"))
            missing = [n for n, v in (("преподаватель", tid), ("группа", gid),
                                      ("аудитория", aid)) if v is None]
            if s and sid is None:
                missing.append("предмет")
            if missing:
                raise ValueError(f"не найдены: {', '.join(missing)}")
            if not 1 <= pair <= 6:
                raise ValueError(f"пара={row.get('Пара')} (нужно 1..6)")
            db.save_schedule(tid, sid, kind, [(gid, aid, pair, weekday)])
            bump(added, "Расписание")
        except Exception as e:
            errors.append(f"Расписание, строка {i}: {e}")

    return {"added": added, "updated": updated, "errors": errors}
